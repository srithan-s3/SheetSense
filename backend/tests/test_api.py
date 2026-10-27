"""HTTP integration tests exercise real parsing, sessions, analysis and exports."""
from io import BytesIO
from pathlib import Path
import csv
import time
import zipfile

import openpyxl
import pytest

from app import create_app


@pytest.fixture
def app(tmp_path):
    return create_app({"TESTING": True, "SESSION_DIR": str(tmp_path / "sessions"), "START_CLEANUP_THREAD": False})


@pytest.fixture
def client(app):
    return app.test_client()


def upload(client, content=b"Category,Amount,Date\nA,10,2026-01-01\nB,30,2026-02-01\nA,20,2026-03-01\n", filename="sales.csv"):
    result = client.post("/api/upload", data={"file": (BytesIO(content), filename)})
    assert result.status_code == 201, result.get_json()
    session = result.get_json()["data"]
    return session, {"X-Session-Token": session["token"]}


def post(client, session, headers, endpoint="analyze", **payload):
    return client.post(f"/api/{endpoint}/{session['session_id']}", headers=headers, json={"sheet": session["sheets"][0], **payload})


def workbook_bytes(sheets):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        worksheet = workbook.create_sheet(name)
        for row in rows:
            worksheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_sales_preview_analysis_and_filter_use_real_data(client):
    session, headers = upload(client)
    analyzed = post(client, session, headers).get_json()["data"]
    assert analyzed["rows"] == 3
    assert analyzed["column_count"] == 3
    assert analyzed["charts"] and analyzed["insights"]
    assert any(chart["type"] == "line" for chart in analyzed["charts"])
    preview = post(client, session, headers, "preview", page=1, page_size=2).get_json()["data"]
    assert len(preview["rows"]) == 2 and preview["total_rows"] == 3
    filtered = post(client, session, headers, "filter", filters={"Category": {"values": ["A"]}}).get_json()["data"]
    assert filtered["rows"] == 2
    assert filtered["kpis"] != analyzed["kpis"]
    searched = post(client, session, headers, "preview", search="30").get_json()["data"]
    assert searched["total_rows"] == 1
    assert float(searched["rows"][0]["Amount"]) == 30


def test_students_multiworksheet_and_empty_worksheet(client):
    content = workbook_bytes({"Marks": [["Student", "Math", "Science"], ["A", 82, 90], ["B", 71, 85]], "Attendance": [["Class", "Attendance"], ["A", 90], ["B", 95]], "Empty": []})
    session, headers = upload(client, content, "students.xlsx")
    assert session["sheets"] == ["Marks", "Attendance", "Empty"]
    assert client.get(f"/api/sheets/{session['session_id']}", headers=headers).get_json()["data"]["sheets"] == session["sheets"]
    assert client.post(f"/api/analyze/{session['session_id']}", headers=headers, json={}).status_code == 400
    response = post(client, session, headers, sheet="Marks")
    assert response.status_code == 200
    analysis = response.get_json()["data"]
    assert analysis["rows"] == 2
    assert "Revenue" not in str(analysis["kpis"])
    assert post(client, session, headers, sheet="Attendance").get_json()["data"]["column_count"] == 2
    empty = post(client, session, headers, sheet="Empty")
    assert empty.status_code == 400 and "empty" in empty.get_json()["error"]["message"].lower()


def test_dates_not_invented_and_categories_without_numbers(client):
    session, headers = upload(client, b"Category,Cost\nFood,12\nTravel,24\nFood,18\n")
    analysis = post(client, session, headers).get_json()["data"]
    assert all(chart["type"] != "line" for chart in analysis["charts"])
    session, headers = upload(client, b"Team,Status\nDesign,Active\nDesign,Active\nSupport,Pending\n")
    analysis = post(client, session, headers).get_json()["data"]
    assert all(column["type"] != "numeric" for column in analysis["columns"])
    assert analysis["charts"]
    assert not any(chart["type"] in ("line", "histogram", "scatter") for chart in analysis["charts"])


def test_quality_cleaning_preview_apply_and_original_preserved(client):
    session, headers = upload(client, b"Category,Amount\nA,10\nA,10\nB,\n")
    analysis = post(client, session, headers).get_json()["data"]
    assert analysis["quality"]["missing_cells"] == 1
    assert analysis["quality"]["duplicate_rows"] == 1
    operations = [{"action": "drop_duplicates"}, {"action": "fill_numeric", "column": "Amount", "method": "mean"}]
    preview = post(client, session, headers, "clean", operations=operations, apply=False).get_json()["data"]
    assert preview["before_rows"] == 3 and preview["after_rows"] == 2 and preview["changed_cells"] == 1
    assert post(client, session, headers, version="cleaned").status_code == 400
    assert post(client, session, headers, "clean", operations=operations, apply=True).status_code == 200
    cleaned = post(client, session, headers, version="cleaned").get_json()["data"]
    assert cleaned["rows"] == 2 and cleaned["quality"]["missing_cells"] == 0
    assert post(client, session, headers).get_json()["data"]["rows"] == 3


def test_exports_reflect_current_filters_and_prevent_formula_injection(client):
    session, headers = upload(client, b"Category,Amount,Note\nA,10,=1+1\nB,30,normal\nA,20,@SUM(A1)\n")
    filters = {"Category": {"values": ["A"]}}
    result = post(client, session, headers, "export", format="csv", filters=filters)
    assert result.status_code == 200 and "attachment" in result.headers["Content-Disposition"]
    rows = list(csv.DictReader(result.data.decode("utf-8-sig").splitlines()))
    assert len(rows) == 2 and {row["Category"] for row in rows} == {"A"}
    assert rows[0]["Note"].startswith("'=") and rows[1]["Note"].startswith("'@")
    result = post(client, session, headers, "export", format="xlsx", filters=filters)
    book = openpyxl.load_workbook(BytesIO(result.data), data_only=False)
    assert book.active.max_row == 3
    assert book.active["C2"].data_type == "s"
    pdf = post(client, session, headers, "export", format="pdf", filters=filters)
    assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF") and len(pdf.data) > 2000
    assert post(client, session, headers, "export", format="exe").status_code == 400


@pytest.mark.parametrize("content,filename", [(b"", "empty.csv"), (b"a,b\n", "headers.csv"), (b"hello", "wrong.exe"), (b"not a workbook", "bad.xlsx"), (b"not xls", "bad.xls"), (b"MZbinary", "bad.csv"), (b"a,b\n1,2,3", "bad.csv"), (b"<html>bad</html>", "bad.csv")])
def test_invalid_uploads_have_safe_errors(client, app, content, filename):
    response = client.post("/api/upload", data={"file": (BytesIO(content), filename)})
    assert response.status_code == 400
    assert response.get_json()["success"] is False
    assert not list(app.extensions["sessions"].root.iterdir())


def test_encoding_delimiters_headers_and_blank_excel_headers(client):
    session, headers = upload(client, "Café;Value;Value;\nParis;1;2;ok\nLyon;3;4;ok\n".encode("cp1252"))
    preview = post(client, session, headers, "preview").get_json()["data"]
    assert [c["name"] for c in preview["columns"]] == ["Café", "Value", "Value (2)", "Column 4"]
    session, headers = upload(client, "Name\tAmount\nA\t3\nB\t4\n".encode("utf-16"))
    assert post(client, session, headers).get_json()["data"]["column_count"] == 2
    session, headers = upload(client, workbook_bytes({"Data": [[None, None], ["a", 2], ["b", 3]]}), "blank-headers.xlsx")
    assert post(client, session, headers).get_json()["data"]["column_count"] == 2


def test_token_isolation_path_safety_deletion_and_expiration(client, app):
    session, headers = upload(client, filename="../../sensitive.csv")
    stored = app.extensions["sessions"].sessions[session["session_id"]]
    assert stored.path.name == "original.csv" and stored.path.is_relative_to(app.extensions["sessions"].root)
    assert session["filename"] == "sensitive.csv"
    assert client.get(f"/api/sheets/{session['session_id']}").status_code == 404
    assert client.get(f"/api/sheets/{session['session_id']}", headers={"X-Session-Token": "wrong"}).status_code == 404
    directory = stored.path.parent
    assert post(client, session, headers, "delete").status_code == 200
    assert not directory.exists()
    assert post(client, session, headers).status_code == 404
    session, headers = upload(client)
    stored = app.extensions["sessions"].sessions[session["session_id"]]
    directory = stored.path.parent
    stored.expires = time.time() - 1
    assert post(client, session, headers).status_code == 404
    assert not directory.exists()


def test_size_row_cells_zip_and_session_limits(client, app):
    app.config["MAX_UPLOAD_BYTES"] = 8
    assert client.post("/api/upload", data={"file": (BytesIO(b"A,B\n1,2\n3,4"), "large.csv")}).status_code == 413
    app.config["MAX_UPLOAD_BYTES"] = 100000
    app.config["MAX_ROWS"] = 1
    assert client.post("/api/upload", data={"file": (BytesIO(b"A,B\n1,2\n3,4"), "rows.csv")}).status_code == 400
    app.config["MAX_ROWS"] = 50000
    app.config["MAX_EXPANDED_BYTES"] = 30
    content = workbook_bytes({"Data": [["A", "B"], [1, 2]]})
    assert client.post("/api/upload", data={"file": (BytesIO(content), "zip.xlsx")}).status_code == 400
    app.config["MAX_SESSIONS"] = 1
    upload(client)
    response = client.post("/api/upload", data={"file": (BytesIO(b"A\n1\n2"), "other.csv")})
    assert response.status_code == 400


def test_invalid_config_filter_and_custom_chart(client):
    session, headers = upload(client)
    assert post(client, session, headers, filters={"not-a-column": {"values": ["A"]}}).status_code == 400
    assert post(client, session, headers, filters=[]).status_code == 400
    assert post(client, session, headers, "custom-chart", x="Category", y="Amount", type="bar", aggregation="sum", sort="desc", top_n=10).status_code == 200
    assert post(client, session, headers, "custom-chart", x="Category", y="bad", type="bar").status_code == 400
    assert post(client, session, headers, "clean", operations=[{"action": "execute", "column": "Amount"}]).status_code == 400


def test_cors_and_consistent_error_envelope(client):
    allowed = client.get("/api/health", headers={"Origin": "http://localhost:5173"})
    assert allowed.headers["Access-Control-Allow-Origin"] == "http://localhost:5173"
    denied = client.get("/api/health", headers={"Origin": "https://untrusted.example"})
    assert "Access-Control-Allow-Origin" not in denied.headers
    assert client.get("/api/missing").get_json()["success"] is False
    assert allowed.headers["Cache-Control"] == "no-store"


def test_excel_formulas_are_never_executed(client):
    content = workbook_bytes({"Data": [["Category", "Amount", "Formula"], ["A", 10, "=1+1"], ["B", 20, "=HYPERLINK(\"https://example.invalid\")"]]})
    session, headers = upload(client, content, "formula.xlsx")
    result = post(client, session, headers, "preview").get_json()["data"]
    assert result["rows"][0]["Formula"] is None


def test_xlsx_false_dimensions_do_not_silently_truncate_data(client):
    original = workbook_bytes({"Data": [["Name", "Amount"], ["A", 3], ["B", 4]]})
    altered = BytesIO()
    with zipfile.ZipFile(BytesIO(original)) as source, zipfile.ZipFile(altered, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = data.replace(b'dimension ref="A1:B3"', b'dimension ref="A1"')
            target.writestr(name, data)
    session, headers = upload(client, altered.getvalue(), "dimensions.xlsx")
    analysis = post(client, session, headers).get_json()["data"]
    assert analysis["rows"] == 2 and analysis["column_count"] == 2
    # Trailing unlabeled columns become generated headers without losing values.
    session, headers = upload(client, workbook_bytes({"Data": [["Name"], ["A", 3], ["B", 4]]}), "trailing.xlsx")
    preview = post(client, session, headers, "preview").get_json()["data"]
    assert preview["rows"][0]["Column 2"] == 3


def test_real_legacy_xls_and_native_dates(client):
    # Golden BIFF8 fixture generated with xlwt; xlwt is not a runtime/test dependency.
    content = (Path(__file__).parent / "fixtures" / "legacy_marks.xls").read_bytes()
    session, headers = upload(client, content, "legacy.xls")
    assert session["sheets"] == ["Legacy marks", "Empty"]
    analysis = post(client, session, headers).get_json()["data"]
    assert analysis["rows"] == 2
    assert next(column for column in analysis["columns"] if column["name"] == "Date")["type"] == "date"
    preview = post(client, session, headers, "preview").get_json()["data"]
    assert preview["rows"][0]["Student"] == "Asha"
    assert preview["rows"][0]["Math"] == 82
    assert post(client, session, headers, sheet="Empty").status_code == 400


def test_unicode_filename_and_invalid_body_types(client):
    session, headers = upload(client, filename="你好.csv")
    assert session["filename"].endswith(".csv")
    for config in ({"sheet": []}, {"version": []}, {"filters": []}):
        assert post(client, session, headers, **config).status_code == 400


def test_extreme_finite_numbers_export_without_invalid_geometry(client):
    session, headers = upload(client, b"Category,Value,Other\nA,-1e308,1e308\nB,1e308,-1e308\nC,0,0\n")
    for format in ("pdf", "csv", "xlsx"):
        response = post(client, session, headers, "export", format=format)
        assert response.status_code == 200, response.get_json()
        assert response.data
    from app.services.exports import chart_drawing
    from reportlab.graphics.shapes import Circle, Line, Rect
    import math
    for kind in ("bar", "horizontal_bar", "line", "scatter", "stacked_bar"):
        drawing = chart_drawing({"type": kind, "x_key": "x", "y_keys": ["y"], "data": [{"x": -1e308, "y": -1e308}, {"x": 1e308, "y": 1e308}]})
        for shape in drawing.contents:
            if isinstance(shape, (Circle, Line, Rect)):
                for name, value in shape.getProperties().items():
                    if name in ("x", "y", "x1", "y1", "x2", "y2", "width", "height", "cx", "cy", "r"):
                        assert math.isfinite(value)


def test_download_samples_use_upload_pipeline(client):
    samples = client.get("/api/samples").get_json()["data"]
    assert {"sales", "students", "expenses", "inventory"}.issubset({sample["id"] for sample in samples})
    for sample in samples:
        downloaded = client.get(f"/api/samples/{sample['id']}")
        assert downloaded.status_code == 200
        session, headers = upload(client, downloaded.data, sample["filename"])
        assert post(client, session, headers).status_code == 200
