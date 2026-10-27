"""Bounded spreadsheet readers. Formulas/macros are never evaluated."""
from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

import pandas as pd
import openpyxl
import xlrd


class InvalidFile(ValueError):
    pass


def _headers(values):
    used, names = set(), []
    for index, value in enumerate(values):
        name = str(value).strip() if value is not None and str(value).strip() else f"Column {index + 1}"
        if len(name) > 200:
            raise InvalidFile("Column names must be 200 characters or fewer.")
        candidate, suffix = name, 2
        while candidate in used:
            candidate = f"{name} ({suffix})"
            suffix += 1
        names.append(candidate)
        used.add(candidate)
    return names


def _frame(rows, config, allow_empty=False, allow_width_growth=False):
    header = next(rows, None)
    if header is None:
        if allow_empty:
            return pd.DataFrame()
        raise InvalidFile("The file is empty. Include column headers and at least one data row.")
    if len(header) > config["MAX_COLUMNS"]:
        raise InvalidFile(f"This dataset exceeds the {config['MAX_COLUMNS']:,}-column limit.")
    names = _headers(header)
    records = []
    for row in rows:
        if not row:
            continue
        if len(row) > len(names):
            if not allow_width_growth:
                raise InvalidFile("A row contains more values than the header. Check the delimiter and column names.")
            if len(row) > config["MAX_COLUMNS"] or (len(records) + 1) * len(row) > config["MAX_CELLS"]:
                raise InvalidFile("Worksheet contents exceed the configured column or cell limits.")
            extra = len(row) - len(names)
            names = _headers(names + [None] * extra)
            for previous in records:
                previous.extend([None] * extra)
        if len(records) >= config["MAX_ROWS"]:
            raise InvalidFile(f"This dataset exceeds the {config['MAX_ROWS']:,}-row limit.")
        if (len(records) + 1) * len(names) > config["MAX_CELLS"]:
            raise InvalidFile(f"This dataset exceeds the {config['MAX_CELLS']:,}-cell limit.")
        if any(isinstance(value, str) and len(value) > 32767 for value in row):
            raise InvalidFile("A cell is too long. The limit is 32,767 characters per cell.")
        values = [None if isinstance(v, str) and not v.strip() else v for v in row]
        records.append(values + [None] * (len(names) - len(values)))
    if not records and not allow_empty:
        raise InvalidFile("This worksheet has no data rows. Choose another sheet or add data below its headers.")
    return pd.DataFrame(records, columns=names, dtype=object)


def _csv(path, config):
    data = path.read_bytes()
    if not data or data.startswith((b"MZ", b"PK", b"%PDF", b"\xd0\xcf\x11\xe0")):
        raise InvalidFile("The file contents do not match a CSV spreadsheet.")
    encodings = ["utf-16"] if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "cp1252"]
    text = None
    for encoding in encodings:
        try:
            text = data.decode(encoding)
            break
        except UnicodeError:
            continue
    if text is None:
        raise InvalidFile("Could not decode this CSV. Save it using UTF-8 or Windows-1252 encoding.")
    if any(ord(c) < 32 and c not in "\t\n\r" for c in text):
        raise InvalidFile("The CSV contains binary content or unsupported control characters.")
    if text.lstrip().lower().startswith(("<!doctype", "<html", "<?xml", "<script", "#!")):
        raise InvalidFile("The file contents do not match a CSV spreadsheet.")
    try:
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        return _frame(iter(csv.reader(io.StringIO(text), dialect=dialect, strict=True)), config)
    except csv.Error as exc:
        raise InvalidFile("The CSV has malformed quoting or rows. Save a fresh CSV and try again.") from exc


def _validate_xlsx(path, config):
    if not zipfile.is_zipfile(path):
        raise InvalidFile("The file is not a valid XLSX workbook.")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(item.file_size for item in entries) > config["MAX_EXPANDED_BYTES"]:
            raise InvalidFile("This workbook expands beyond the safe processing limit. Use a smaller workbook.")
        names = {item.filename for item in entries}
        if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
            raise InvalidFile("The file contents do not match an XLSX workbook.")
        if any("vbaproject" in name.lower() or name.startswith("xl/embeddings/") for name in names):
            raise InvalidFile("Workbooks with macros or embedded objects are unsupported. Save a plain XLSX file.")
        if any(item.flag_bits & 1 for item in entries):
            raise InvalidFile("Password-protected workbooks are unsupported.")


def inspect_file(path: Path, config):
    """Validate every sheet under cumulative resource bounds before accepting."""
    extension = path.suffix.lower()
    if extension == ".csv":
        _csv(path, config)
        return ["Data"]
    if extension == ".xlsx":
        _validate_xlsx(path, config)
    elif extension == ".xls":
        if path.read_bytes()[:8] != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise InvalidFile("The file contents do not match a legacy XLS workbook.")
    else:
        raise InvalidFile("Choose an .xlsx, .xls, or .csv file.")
    try:
        if extension == ".xlsx":
            with _open_xlsx(path) as workbook:
                names = workbook.sheetnames
                if len(names) > config["MAX_SHEETS"]:
                    raise InvalidFile(f"Workbooks may contain at most {config['MAX_SHEETS']} worksheets.")
                total = 0
                for name in names:
                    frame = _xlsx_frame(workbook[name], config)
                    total += frame.size
                    if total > config["MAX_CELLS"]:
                        raise InvalidFile("The workbook contains too many cells across its worksheets.")
                return names
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            if workbook.nsheets > config["MAX_SHEETS"]:
                raise InvalidFile(f"Workbooks may contain at most {config['MAX_SHEETS']} worksheets.")
            total = 0
            for name in workbook.sheet_names():
                frame = _xls_frame(workbook, name, config)
                total += frame.size
                if total > config["MAX_CELLS"]:
                    raise InvalidFile("The workbook contains too many cells across its worksheets.")
            return workbook.sheet_names()
        finally:
            workbook.release_resources()
    except InvalidFile:
        raise
    except Exception as exc:
        raise InvalidFile("This workbook could not be read. Check that it is valid and not password-protected.") from exc


class _open_xlsx:
    def __init__(self, path):
        self.path = path

    def __enter__(self):
        self.book = openpyxl.load_workbook(self.path, read_only=True, data_only=True, keep_links=False)
        return self.book

    def __exit__(self, *args):
        self.book.close()


def _xlsx_frame(sheet, config):
    if (sheet.max_row or 0) > config["MAX_ROWS"] + 1 or (sheet.max_column or 0) > config["MAX_COLUMNS"]:
        raise InvalidFile("Worksheet dimensions exceed the configured row or column limits. Remove unused formatted rows or columns.")
    # Do not trust producer-supplied dimensions to truncate actual cells.
    # Both row and column growth are checked while consuming the stream.
    sheet.reset_dimensions()
    rows = iter(sheet.iter_rows(values_only=True))
    first = next(rows, None)
    if first is None:
        return pd.DataFrame()
    from itertools import chain
    frame = _frame(iter(chain([first], rows)), config, allow_empty=True, allow_width_growth=True)
    return frame.dropna(how="all").reset_index(drop=True)


def _xls_frame(workbook, name, config):
    sheet = workbook.sheet_by_name(name)
    if sheet.nrows > config["MAX_ROWS"] + 1 or sheet.ncols > config["MAX_COLUMNS"]:
        raise InvalidFile("Worksheet dimensions exceed the configured row or column limits.")
    def rows():
        for row_index in range(sheet.nrows):
            values = []
            for cell in sheet.row(row_index):
                if cell.ctype == xlrd.XL_CELL_DATE:
                    values.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    values.append(None)
                else:
                    values.append(cell.value)
            yield values
    return _frame(iter(rows()), config, allow_empty=True)


def read_sheet(path: Path, sheet: str, config):
    try:
        if path.suffix == ".csv":
            return _csv(path, config)
        if path.suffix == ".xlsx":
            with _open_xlsx(path) as workbook:
                frame = _xlsx_frame(workbook[sheet], config)
        else:
            workbook = xlrd.open_workbook(str(path), on_demand=True)
            try:
                frame = _xls_frame(workbook, sheet, config)
            finally:
                workbook.release_resources()
        if frame.empty:
            raise InvalidFile("This worksheet is empty. Choose another worksheet to continue.")
        return frame
    except InvalidFile:
        raise
    except Exception as exc:
        raise InvalidFile("The selected worksheet could not be read.") from exc
